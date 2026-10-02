import os
import shutil
import tempfile
import zipfile

from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook

from src.file_utils import clean_filename


def load_workbook_data(file_path):
    """
    Carrega o workbook para leitura da interface.
    """

    return load_workbook(
        file_path,
        data_only=False,
        keep_vba=False
    )


def get_sheet_names(workbook):
    return workbook.sheetnames


def get_columns(
    workbook,
    sheet_name,
    return_unique=False,
    filter_column=None
):
    """
    Retorna os cabeçalhos ou os valores únicos da coluna selecionada.
    """

    sheet = workbook[sheet_name]

    headers = [
        cell.value
        for cell in sheet[1]
    ]

    if not return_unique:
        return headers

    column_index = headers.index(filter_column)

    unique_values = set()

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):
        value = row[column_index]

        if value is not None:
            unique_values.add(str(value))

    return sorted(unique_values)


def remove_excel_tables(sheet):
    """
    Remove todas as estruturas Table/ListObject
    mantendo a formatação visual da planilha.
    """

    table_names = list(sheet.tables.keys())

    for table_name in table_names:
        del sheet.tables[table_name]


def keep_only_selected_sheet(
    workbook,
    selected_sheet
):
    """
    Remove todas as abas, mantendo somente a selecionada.
    """

    for sheet_name in list(workbook.sheetnames):

        if sheet_name != selected_sheet:
            workbook.remove(
                workbook[sheet_name]
            )


def get_filter_column_index(sheet, filter_column):
    """
    Localiza o índice da coluna de filtro.
    """

    for cell in sheet[1]:

        if cell.value == filter_column:
            return cell.column

    raise ValueError(
        f"A coluna '{filter_column}' não foi encontrada."
    )


def remove_non_matching_rows(
    sheet,
    filter_column,
    filter_value
):
    """
    Remove fisicamente as linhas que não correspondem
    ao valor selecionado.

    A exclusão é feita em blocos contínuos para reduzir
    drasticamente o número de operações do OpenPyXL.
    """

    column_index = get_filter_column_index(
        sheet,
        filter_column
    )

    rows_to_delete = []

    for row_number in range(
        2,
        sheet.max_row + 1
    ):

        value = sheet.cell(
            row=row_number,
            column=column_index
        ).value

        if str(value) != str(filter_value):
            rows_to_delete.append(row_number)

    if not rows_to_delete:
        return

    # Agrupa linhas consecutivas.
    blocks = []

    start = rows_to_delete[0]
    previous = rows_to_delete[0]

    for row in rows_to_delete[1:]:

        if row == previous + 1:
            previous = row

        else:
            blocks.append(
                (start, previous)
            )

            start = row
            previous = row

    blocks.append(
        (start, previous)
    )

    # Exclui de baixo para cima.
    for start, end in reversed(blocks):

        sheet.delete_rows(
            start,
            end - start + 1
        )


def generate_single_file(
    source_file,
    sheet_name,
    filter_column,
    filter_value,
    output_dir
):
    """
    Gera um único arquivo filtrado.

    Cada arquivo é aberto diretamente a partir da
    fonte original, evitando deepcopy do Workbook.
    """

    workbook = load_workbook(
        source_file,
        data_only=True,
        keep_vba=False
    )

    try:

        keep_only_selected_sheet(
            workbook,
            sheet_name
        )

        sheet = workbook[sheet_name]

        remove_non_matching_rows(
            sheet,
            filter_column,
            filter_value
        )

        remove_excel_tables(sheet)

        current_date = datetime.now()

        month = current_date.strftime("%m")
        year = current_date.strftime("%Y")

        safe_value = clean_filename(
            str(filter_value)
        )

        filename = (
            f"{filter_column}_"
            f"{safe_value}_"
            f"{month}_"
            f"{year}.xlsx"
        )

        final_path = output_dir / filename

        workbook.save(final_path)

        return final_path

    finally:

        workbook.close()


def generate_filtered_files(
    source_file,
    sheet_name,
    filter_column,
    values,
    progress_bar,
    status_text
):
    """
    Gera todos os arquivos filtrados e cria um ZIP final.

    Versão otimizada:
    - não utiliza deepcopy
    - não cria cópia intermediária do Excel
    - exclui linhas em blocos
    - utiliza diretório temporário exclusivo
    - gera ZIP somente com os arquivos desta execução
    """

    output_dir = Path(
        tempfile.mkdtemp(
            prefix="excel_filter_"
        )
    )

    generated_files = []

    total = len(values)

    try:

        for index, value in enumerate(values):

            status_text.text(
                f"Gerando {index + 1} de {total}: {value}"
            )

            file_path = generate_single_file(
                source_file=source_file,
                sheet_name=sheet_name,
                filter_column=filter_column,
                filter_value=value,
                output_dir=output_dir
            )

            generated_files.append(
                file_path
            )

            progress = int(
                ((index + 1) / total) * 100
            )

            progress_bar.progress(
                progress
            )

        zip_path = output_dir / "arquivos_filtrados.zip"

        with zipfile.ZipFile(
            zip_path,
            "w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=6
        ) as zip_file:

            for file_path in generated_files:

                zip_file.write(
                    file_path,
                    arcname=file_path.name
                )

        status_text.text(
            "Processo finalizado!"
        )

        return zip_path

    except Exception:

        # Limpa arquivos caso ocorra erro.
        for file_path in output_dir.glob("*"):

            try:
                file_path.unlink()
            except Exception:
                pass

        try:
            output_dir.rmdir()
        except Exception:
            pass

        raise
